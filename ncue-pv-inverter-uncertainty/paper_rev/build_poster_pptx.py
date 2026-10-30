"""CSME 2026 poster (A0 portrait) built inside the official conference template.
Usage: python build_poster_pptx.py <Poster_Template.pptx> <out.pptx>
Figures: ../poster_figs/*.png (900 dpi re-renders of the paper figures)."""
import sys, copy, os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

TPL, OUT = sys.argv[1], sys.argv[2]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "poster_figs")
prs = Presentation(TPL)
s = prs.slides[0]
NAVY = RGBColor(0x1A, 0x3E, 0x8C)
ORANGE = RGBColor(0xB8, 0x5A, 0x12)
INK = RGBColor(0x1B, 0x1F, 0x24)
FONT = "Arial"

# ------------------------------------------------------------------ template header fields
shapes = {sh.name: sh for sh in s.shapes}
title_sh, info_sh = [sh for sh in s.shapes if sh.has_text_frame][:2]
paperno_sh = [sh for sh in s.shapes if sh.has_text_frame and "論文編號" in sh.text_frame.text][0]


def set_para_text(p, text, size=None, sup=None):
    """Keep the first run's formatting, drop the others."""
    runs = p.runs
    r0 = runs[0]
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = text
    if size:
        r0.font.size = Pt(size)
    return r0


tp = title_sh.text_frame.paragraphs[0]
set_para_text(tp, "A pvlib-Based Computational Model of a Multi-Inverter Rooftop Photovoltaic Plant Driven by "
              "Public Weather Data: Sensitivity Analysis and Field Validation", size=60)
title_sh.top = Cm(5.0)

ps = info_sh.text_frame.paragraphs
r = set_para_text(ps[0], "Cliff Nathanael Govino*, Wen-Ren Yang")
set_para_text(ps[1], "International Master Program of Engineering, National Changhua University of Education, Changhua, Taiwan")
set_para_text(ps[3], "*E-mail: cliffgovin@gmail.com")
ps[2]._p.getparent().remove(ps[2]._p)          # NSTC grant line: no grant
info_sh.left, info_sh.width, info_sh.top = Cm(2.0), Cm(64.0), Cm(12.2)
for para in info_sh.text_frame.paragraphs:
    for rr in para.runs:
        rr.font.size = Pt(34)
info_sh.top = Cm(11.3)
set_para_text(paperno_sh.text_frame.paragraphs[0], "論文編號：A10-006")
paperno_sh.top = Cm(15.6)

# ------------------------------------------------------------------ helpers
X0, X1 = 1.6, 82.5
GAP = 2.2
COLW = (X1 - X0 - 2 * GAP) / 3
COLX = [X0 + i * (COLW + GAP) for i in range(3)]


def text(x, y, w, h, paras, size=28, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.2)
    tf.margin_top = tf.margin_bottom = Cm(0.1)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.35)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            rr = p.add_run()
            rr.text = t.replace(" %", "\u00a0%")
            f = rr.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
            if o.get("sub"):
                rr._r.get_or_add_rPr().set("baseline", "-25000")
    return tb


def heading(x, y, w, num, label):
    text(x, y, w, 1.9, [[(f"{num}  " if num else "", {"color": ORANGE, "bold": True}), (label, {"bold": True})]],
         size=40, color=NAVY, name=f"Heading {label}")
    ln = s.shapes.add_connector(1, Cm(x + 0.2), Cm(y + 2.0), Cm(x + w - 0.2), Cm(y + 2.0))
    ln.line.color.rgb = NAVY
    ln.line.width = Pt(2.5)
    return y + 2.5


def picture(name, x, y, w, alt):
    path = os.path.join(FIG, name)
    W, H = Image.open(path).size
    h = w * H / W
    pic = s.shapes.add_picture(path, Cm(x), Cm(y), Cm(w), Cm(h))
    pic.name = alt
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return y + h


def box(x, y, w, h, fill):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    b.adjustments[0] = 0.06
    b.fill.solid()
    b.fill.fore_color.rgb = fill
    b.line.fill.background()
    b.shadow.inherit = False
    return b


# ------------------------------------------------------------------ key results band
Y = 20.2
stats = [("≤ 3 %", NAVY, RGBColor(0xEA, 0xF0, 0xFB),
          "Over 3.5 years the inverter energy counters agree with the model (1411 and 1437 vs 1399 kWh/kWp per year)."),
         ("81 %", ORANGE, RGBColor(0xFB, 0xEF, 0xE3),
          "On clear autumn days the plant delivers only 81 % of the modeled energy; no plausible set of factor values explains the gap."),
         ("2.6 % vs 0.5 %", NAVY, RGBColor(0xEA, 0xF0, 0xFB),
          "The inverter-model effect comes from the assumed efficiency, not from the model structure.")]
for (big, col, fill, small), x in zip(stats, COLX):
    box(x, Y, COLW, 7.4, fill)
    text(x + 0.6, Y + 0.4, COLW - 1.2, 3.2, [[(big, {"bold": True, "color": col})]], size=80, name=f"Stat {big}")
    text(x + 0.6, Y + 3.9, COLW - 1.2, 4.0, [small], size=26)

# ------------------------------------------------------------------ row A
YA = 29.0
y = heading(COLX[0], YA, 2 * COLW + GAP, "1", "Plant and measurement point")
y = picture("fig1_layout.png", COLX[0] + (2 * COLW + GAP - 46.0) / 2, y + 0.2, 46.0,
            "Roof blocks, strings per MPPT, three inverters and the Hioki recorder at the common AC output")
text(COLX[0], y + 0.3, 2 * COLW + GAP, 2.2,
     ["372 modules (139.5 kWp) at 23° tilt, facing south; as-built strings per MPPT from the inverter displays. "
      "The recorder measures the sum of the three inverters at the AC panel."], size=26)
rowA_end = y + 2.6

yq = heading(COLX[2], YA, COLW, "", "Research questions")
text(COLX[2], yq, COLW, 11.5, [
    [("(i)  ", {"bold": True, "color": ORANGE}), ("How large is the inverter-model effect compared with other uncertain factors, and is it the model structure or the assumed efficiency?", {})],
    [("(ii)  ", {"bold": True, "color": ORANGE}), ("Does it matter whether irradiance comes from a station 21 km away or from ERA5?", {})],
    [("(iii)  ", {"bold": True, "color": ORANGE}), ("Does the uncertainty budget enclose the energy the plant really delivers?", {})]], size=28)
yd = heading(COLX[2], yq + 11.6, COLW, "", "Data and model")
text(COLX[2], yd, COLW, 13.0, [
    [("Plant: ", {"bold": True}), ("2 × 50 kW + 30 kW inverters with 145 / 140 / 87 modules", {})],
    [("Weather: ", {"bold": True}), ("CWA Tianzhong station (hourly, 21.4 km) and ERA5, 2023", {})],
    [("Model (pvlib): ", {"bold": True}), ("Erbs → Perez → Ross → PVWatts DC (8.7 % loss) → 3 inverters", {})],
    [("Measured: ", {"bold": True}), ("Hioki MR8875, 16–19 Oct 2026 and seven 2-h records; inverter lifetime counters", {})]], size=28)

# ------------------------------------------------------------------ row B (three columns)
YB = max(rowA_end, yd + 12.0) + 1.2


def section(col, y, num, label, fig, fw, alt, body, th):
    x = COLX[col]
    y = heading(x, y, COLW, num, label)
    y = picture(fig, x + (COLW - fw) / 2, y + 0.2, fw, alt)
    text(x, y + 0.3, COLW, th, body, size=26)
    return y + 0.3 + th + 0.9


sub = lambda t: (t, {"sub": True, "size": 26})
it = lambda t: (t, {"italic": True})
# column 1
y = section(0, YB, "2", "Three-phase power", "fig2_3phase_wave.png", 18.0,
            "Three phase voltages 120 degrees apart, rms and peak values",
            [[("Bus voltages: three sinusoids 120° apart. Balanced: ", {}), it("P"), (" = √3 ", {}), it("V"), sub("LL"), (" ", {}), it("I"), sub("L"), (" cos φ. "
              "With three wires the recorder uses two products, ", {}), it("P"), (" = ⟨", {}), it("v"), sub("ab"), it("i"), sub("a"), (" − ", {}), it("v"), sub("bc"), it("i"), sub("c"), ("⟩; "
              "thinning 2 ms → 10 ms changes the power by ≤ 0.12 %.", {})]], 5.4)
y = section(0, y, "3", "Annual uncertainty budget", "fig_oat.png", 21.0,
            "Change in annual energy per factor and Sobol total index",
            ["Sobol, nine factors, 45,056 runs (converged). DC loss and GHI error explain 84 % of the variance, the inverter model 7 %. "
             "Baseline 195 MWh (1399 kWh/kWp); 5–95 %: −5.5 % to +11.4 %."], 5.4)
# column 2
y = section(1, YB, "4", "Inverter: structure or efficiency?", "fig_inv.png", 22.0,
            "Annual energy for constant and PVWatts inverter at equal efficiency",
            ["Same η: constant efficiency and the PVWatts curve differ by only 0.4–0.6 %. "
             "Lowering η from 0.985 to 0.96 removes 2.6 %: the inverter share of the budget is an efficiency assumption."], 5.4)
y = section(1, y, "5", "Model vs. measurement", "fig_modelvs_rev.png", 15.5,
            "Hourly measured and modeled AC power on 17 and 18 October 2026",
            [[it("R"), (" = Σ", {}), it("P"), sub("actual"), (" / Σ", {}), it("P"), sub("model"), (". The hourly shape is reproduced (", {}), it("r"),
              (" = 0.99, 0.97), but the model is higher in almost every hour: ", {}), it("R"), (" = 0.81 and 0.80 (nMBE +23 %, +24 %).", {})]], 5.4)
# column 3
y = section(2, YB, "6", "Closing the budget", "fig_closure_rev.png", 19.5,
            "Modeled energy with all lower bounds stacked against the measurement",
            ["Every factor at its lowest end (incl. isotropic transposition and Boland decomposition) gives 82 %, close to the measured 81 %. "
             "This joint extreme lowers the annual energy by 13.6 % (5th percentile −5.5 %) and underestimates 10 Feb by 6 %: implausible."], 6.6)
y = section(2, y, "7", "Long-term check", "fig_longterm.png", 20.0,
            "Measured over modeled energy for clear days and 3.5-year inverter counters",
            ["Lifetime counters of the two 50 kW inverters (1,280 days) are 0.8 % and 2.7 % above the model; the 30 kW counter is faulty. "
             "The clear-day shortfall is not a permanent loss."], 4.6)
rowB_end = y

# ------------------------------------------------------------------ conclusion band
YC = 107.2
box(X0, YC, X1 - X0, 6.6, RGBColor(0xEA, 0xF0, 0xFB))
cw = (X1 - X0 - 2 * GAP) / 3
items = [("Conclusion. ", "Suitable for annual energy, but not for hourly prediction on clear days (about 20 % too high)."),
         ("Next step. ", "Irradiance sensor on the roof and monthly utility-meter data to separate irradiance and plant losses."),
         ("Code and data. ", "github.com/cliff1221/ncue-pv-inverter")]
for (lab, t), x in zip(items, COLX):
    text(x + 0.5, YC + 0.3, cw - 1.0, 6.0, [[(lab, {"bold": True, "color": NAVY}), (t, {})]], size=28, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT, "rowA_end %.1f YB %.1f rowB_end %.1f (conclusion at %.1f)" % (rowA_end, YB, rowB_end, YC))
