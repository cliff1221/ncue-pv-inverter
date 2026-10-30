"""Move the final revised paper into the official CSME 2026 Word template.

Keeps the template's page setup, header, styles, two-column body and reference numbering,
and copies the text (with italic/bold/sub/superscript), equations, figures and tables
from CSME2026_Govino_revised.docx.
Usage: python to_template.py <template.docx> <out.docx>
"""
import sys, io, copy, re
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Emu, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.text.paragraph import Paragraph

SRC = "CSME2026_Govino_revised.docx"
TPL, OUT = sys.argv[1], sys.argv[2]
src, doc = Document(SRC), Document(TPL)
body = doc.element.body
COLW = (11905 - 1134 - 1134 - 559) // 2          # twips, one column of the template
MAXW_IN = 2.85
# per-figure width (in): smaller for simple charts
FIGW = {2: 2.7, 3: 2.85, 4: 2.85, 5: 2.6, 6: 2.85, 7: 2.4, 8: 2.55, 9: 2.4, 10: 2.35, 11: 2.6, 12: 2.6}

COLOR = '<w:color w:val="000000" w:themeColor="text1"/>'


def rpr(sz=20, b=False, i=False, va=None, extra=""):
    s = "<w:rPr>" + ('<w:b/><w:bCs/>' if b else '') + ('<w:i/><w:iCs/>' if i else '') + COLOR + extra
    s += f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/>'
    if va:
        s += f'<w:vertAlign w:val="{va}"/>'
    return s + "</w:rPr>"


def run_xml(text, **k):
    parts = text.split("\t")
    out = ""
    for j, part in enumerate(parts):
        if j:
            out += f"<w:r>{rpr(**k)}<w:tab/></w:r>"
        if part:
            out += f'<w:r>{rpr(**k)}<w:t xml:space="preserve">{escape(part)}</w:t></w:r>'
    return out


def runs_from(p, sz=20, force_b=None):
    out = ""
    for r in p.runs:
        va = "subscript" if r.font.subscript else ("superscript" if r.font.superscript else None)
        b = r.bold if force_b is None else force_b
        out += run_xml(r.text, sz=sz, b=bool(b), i=bool(r.italic), va=va)
    return out


def ppr(jc=None, ind=None, before=None, after=None, keep=False, tabs=None, extra=""):
    s = '<w:pPr><w:pStyle w:val="Default"/>' + (tabs or "")
    if keep:
        s += "<w:keepNext/>"
    if before is not None or after is not None:
        s += f'<w:spacing w:before="{before or 0}" w:after="{after or 0}"/>'
    if ind:
        s += ind
    if jc:
        s += f'<w:jc w:val="{jc}"/>'
    return s + extra + "</w:pPr>"


SECT_END = body[-1]                       # body sectPr (two columns)


def append_xml(xml):
    el = parse_xml(f"<w:p {nsdecls('w')}>{xml}</w:p>")
    SECT_END.addprevious(el)
    return el


def append_image(src_p, width_in):
    blip = src_p._p.find(".//" + qn("a:blip"))
    rid = blip.get(qn("r:embed"))
    blob = src.part.related_parts[rid].blob
    ext = src_p._p.find(".//" + qn("wp:extent"))
    w = min(Emu(int(ext.get("cx"))).inches, width_in)
    p = doc.add_paragraph()
    p.alignment = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = 0
    p.add_run().add_picture(io.BytesIO(blob), width=Inches(w))
    return p


# ---------------------------------------------------------------- title block (section 1)
kids = list(body.iterchildren())
sect1 = next(i for i, el in enumerate(kids) if el.tag == qn("w:p") and el.find(".//" + qn("w:sectPr")) is not None)
for el in kids[sect1 + 1:-1]:
    body.remove(el)
tp = [Paragraph(el, doc) for el in kids[:sect1]]


def set_runs(p, xml_runs):
    for r in list(p._p.iterchildren(qn("w:r"))):
        p._p.remove(r)
    for r in parse_xml(f"<w:x {nsdecls('w')}>{xml_runs}</w:x>").iterchildren():
        p._p.append(r)


title = src.paragraphs[0].text
set_runs(tp[0], run_xml(title, sz=28, b=True))
set_runs(tp[2], run_xml("Cliff Nathanael Govino", sz=24) + run_xml("*", sz=24, va="superscript") + run_xml(", Wen-Ren Yang", sz=24))
set_runs(tp[3], run_xml("International Master Program of Engineering, National Changhua University of Education, Changhua, Taiwan", sz=24))
for p in (tp[4], tp[5]):                  # second affiliation and NSTC grant line: not applicable
    p._p.getparent().remove(p._p)
set_runs(tp[6], run_xml("*", sz=24, va="superscript") + run_xml("Email: cliffgovin@gmail.com", sz=24))

# Fig. 1 full width in the one-column title section
sps = src.paragraphs
fig1_img = next(p for p in sps if p._p.find(".//" + qn("w:drawing")) is not None)
fig1_cap = sps[sps.index(fig1_img) + 1]
sect_p = kids[sect1]
pi = append_image(fig1_img, 6.0)
sect_p.addprevious(pi._p)
cap = append_xml(ppr(jc="both", after=60) + runs_from(fig1_cap, sz=20))
sect_p.addprevious(cap)

# ---------------------------------------------------------------- body (section 2, two columns)
src_kids = list(src.element.body.iterchildren())
start = next(i for i, el in enumerate(src_kids) if el.tag == qn("w:p") and Paragraph(el, src).text == "Abstract")
eq_tabs = f'<w:tabs><w:tab w:val="center" w:pos="{COLW // 2}"/><w:tab w:val="right" w:pos="{COLW}"/></w:tabs>'
renum = {"Data Availability": "7. Data Availability", "References": "8. References"}
for el in src_kids[start:]:
    tag = el.tag.split("}")[1]
    if tag == "sectPr":
        continue
    if tag == "tbl":
        t = copy.deepcopy(el)
        SECT_END.addprevious(t)
        continue
    p = Paragraph(el, src)
    text = p.text
    has_img = el.find(".//" + qn("w:drawing")) is not None
    sz = next((r.font.size.pt for r in p.runs if r.font.size), None)
    if has_img:
        nxt = Paragraph(src_kids[src_kids.index(el) + 1], src).text
        fn = re.match(r"Fig\. (\d+)", nxt)
        append_image(p, FIGW.get(int(fn.group(1)) if fn else 0, MAXW_IN))
    elif not text.strip():
        continue
    elif sz == 11 and p.alignment == 1:                                  # section heading
        t = renum.get(text, text)
        append_xml(ppr(jc="center", before=120, after=40, keep=True) + run_xml(t, sz=22, b=True))
    elif re.match(r"^\d\.\d ", text):                                    # subsection
        append_xml(ppr(before=120, after=40, keep=True) + run_xml(text, sz=20, b=True))
    elif text.startswith("Keywords"):
        append_xml(ppr(jc="both", before=60) + run_xml("Keywords", sz=20, i=True) + run_xml(text[len("Keywords"):], sz=20))
    elif text.startswith("\t"):                                          # equation
        append_xml(ppr(tabs=eq_tabs, before=60, after=60) + runs_from(p, sz=20))
    elif re.match(r"^\d+\. ", text):                                # reference
        num = '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="9"/></w:numPr>'
        xml = ppr(jc="both", ind='<w:ind w:left="340" w:hanging="340"/>', extra="")
        xml = xml.replace('<w:pStyle w:val="Default"/>', '<w:pStyle w:val="Default"/>' + num)
        append_xml(xml + run_xml(re.sub(r"^\d+\. ", "", text), sz=20))
    elif sz == 9 and text.startswith("Table"):                           # table caption
        append_xml(ppr(jc="center", before=80, after=40, keep=True) + runs_from(p, sz=20))
    elif sz == 9 and text.startswith("Fig"):                             # figure caption
        append_xml(ppr(jc="both", after=60) + runs_from(p, sz=20))
    else:                                                                # body text
        indent = p.paragraph_format.first_line_indent is not None or "w:firstLine" in el.xml
        ind = '<w:ind w:firstLine="510"/>' if indent else None
        keep = text.rstrip().endswith((":", "ratio", "is", "[3],"))
        append_xml(ppr(jc="both", ind=ind, keep=keep) + runs_from(p, sz=20))

doc.save(OUT)
print("saved", OUT)
